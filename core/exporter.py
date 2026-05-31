"""Excel 导出模块"""
import os
import sys
from datetime import datetime
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def export_to_excel(comments: List[dict], filepath: str = None) -> str:
    """
    将评论数据导出为 Excel 文件
    返回文件路径
    """
    if filepath is None:
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        # 写入数据目录（兼容 PyInstaller 打包）
        if getattr(sys, 'frozen', False):
            base = os.path.dirname(sys.executable)
        else:
            base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # 项目根目录
        out_dir = os.path.join(base, 'data', 'exports')
        os.makedirs(out_dir, exist_ok=True)
        filepath = os.path.join(out_dir, f'团建线索_{timestamp}.xlsx')

    wb = Workbook()
    ws = wb.active
    ws.title = '销售线索'

    # 表头
    headers = ['序号', '意向', '评论内容', '评论者昵称', '评论者ID', '评论时间',
               '笔记标题', '笔记链接', '搜索关键词', '采集时间']
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(name='微软雅黑', bold=True, color='FFFFFF', size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    # 数据行
    intent_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    normal_font = Font(name='微软雅黑', size=10)
    link_font = Font(name='微软雅黑', size=10, color='0563C1', underline='single')

    for i, c in enumerate(comments, 1):
        row = i + 1
        values = [
            i,
            '🔥 意向' if c.get('is_intent') else '',
            c.get('content', ''),
            c.get('user_name', ''),
            c.get('user_id', ''),
            c.get('created_at', ''),
            c.get('note_title', ''),
            c.get('note_url', ''),
            c.get('keyword', ''),
            c.get('collected_at', ''),
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = normal_font
            cell.alignment = Alignment(vertical='center', wrap_text=(col == 3))
            cell.border = thin_border
            if c.get('is_intent'):
                cell.fill = intent_fill

        # 笔记链接列设为超链接
        note_url = c.get('note_url', '')
        if note_url:
            link_cell = ws.cell(row=row, column=8)
            link_cell.hyperlink = note_url
            link_cell.font = link_font

    # 列宽
    col_widths = [6, 8, 50, 16, 16, 18, 30, 40, 14, 18]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    # 冻结首行
    ws.freeze_panes = 'A2'
    # 筛选
    ws.auto_filter.ref = ws.dimensions

    wb.save(filepath)
    return filepath
